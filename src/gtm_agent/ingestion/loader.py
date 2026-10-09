from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from gtm_agent.models.schemas import ProductBrief


SUPPORTED_SUFFIXES = {
    ".txt": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".pdf": "pdf",
    ".csv": "spreadsheet",
    ".xlsx": "spreadsheet",
    ".zip": "notion_export",
}


def normalize_text(text: str) -> str:
    """Normalize line endings and whitespace while retaining paragraph boundaries."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [re.sub(r"[ \t]+", " ", part).strip() for part in re.split(r"\n\s*\n", text)]
    return "\n\n".join(part for part in paragraphs if part)


def _csv_text(data: bytes) -> str:
    decoded = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(decoded)))
    if not rows:
        return ""
    headers = [cell.strip() or f"Column {index + 1}" for index, cell in enumerate(rows[0])]
    records: list[str] = []
    for row_number, row in enumerate(rows[1:], start=2):
        fields = [f"{headers[index]}: {value.strip()}" for index, value in enumerate(row) if value.strip()]
        if fields:
            records.append(f"Row {row_number}\n" + "\n".join(fields))
    return "\n\n".join(records)


def _xlsx_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else f"Column {i + 1}" for i, value in enumerate(rows[0])]
        records: list[str] = []
        for row_number, row in enumerate(rows[1:], start=2):
            fields = [
                f"{headers[index]}: {value}"
                for index, value in enumerate(row)
                if value is not None and str(value).strip()
            ]
            if fields:
                records.append(f"Row {row_number}\n" + "\n".join(fields))
        if records:
            sheets.append(f"Sheet: {sheet.title}\n\n" + "\n\n".join(records))
    return "\n\n".join(sheets)


def _single_document(filename: str, data: bytes, *, source_id: str | None = None) -> ProductBrief:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES or suffix == ".zip":
        raise ValueError(f"unsupported document type: {suffix or '<none>'}")
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        raw = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".csv":
        raw = _csv_text(data)
    elif suffix == ".xlsx":
        raw = _xlsx_text(data)
    else:
        raw = data.decode("utf-8-sig")
    content = normalize_text(raw)
    if not content:
        raise ValueError(f"document contains no readable content: {filename}")
    identifier = source_id or Path(filename).name
    title = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    if suffix in {".md", ".markdown"}:
        heading = next((line.lstrip("#").strip() for line in content.splitlines() if line.startswith("# ")), "")
        title = heading or title
    return ProductBrief(
        title=title,
        source_id=identifier,
        content=content,
        media_type=SUPPORTED_SUFFIXES[suffix],
    )


def load_bytes(filename: str, data: bytes) -> list[ProductBrief]:
    """Load an uploaded file, expanding a Notion ZIP into attributed documents."""
    suffix = Path(filename).suffix.lower()
    if suffix != ".zip":
        return [_single_document(filename, data)]
    documents: list[ProductBrief] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for member in sorted(archive.infolist(), key=lambda item: item.filename.casefold()):
            member_suffix = Path(member.filename).suffix.lower()
            if member.is_dir() or member_suffix not in {".txt", ".md", ".markdown", ".csv"}:
                continue
            documents.append(
                _single_document(
                    member.filename,
                    archive.read(member),
                    source_id=f"{Path(filename).name}:{member.filename}",
                )
            )
    if not documents:
        raise ValueError(f"Notion export contains no supported documents: {filename}")
    return documents


def load_document(path: str | Path) -> ProductBrief:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    documents = load_bytes(source.name, source.read_bytes())
    return documents[0] if len(documents) == 1 else merge_documents(documents)


def load_documents(paths: list[str | Path]) -> list[ProductBrief]:
    documents: list[ProductBrief] = []
    for path in paths:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        documents.extend(load_bytes(source.name, source.read_bytes()))
    if not documents:
        raise ValueError("at least one readable source document is required")
    return documents


def merge_documents(documents: list[ProductBrief]) -> ProductBrief:
    if not documents:
        raise ValueError("at least one document is required")
    content = "\n\n".join(f"[Source: {item.source_id}]\n{item.content}" for item in documents)
    return ProductBrief(
        title=documents[0].title,
        source_id=documents[0].source_id if len(documents) == 1 else f"collection:{len(documents)}",
        content=content,
        media_type="text",
    )
